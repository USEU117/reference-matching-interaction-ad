from pathlib import Path
import copy, re, json, csv, hashlib, zipfile
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).resolve().parent
OUT=ROOT/'docs/paper_complete_review_20260920'
OUT.mkdir(parents=True,exist_ok=True)
REF=ROOT/'docs/manuscript_polished_20260919/Reference_Matching_English_Polished_20260919.docx'
SOURCE_SHA=hashlib.sha256(REF.read_bytes()).hexdigest()
d=Document(REF)
# Reuse the retained manuscript package, styles, page system, and page footer.
# All old-topic body content is an explicitly replaceable slot.
for e in list(d.element.body):
    if e.tag!=qn('w:sectPr'):d.element.body.remove(e)
for e in list(d.styles.element.iter(qn('w:pBdr')))+list(d.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
for n in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
    d.styles[n].font.color.rgb=RGBColor(0,0,0)
    d.styles[n].font.name='Times New Roman'
d.core_properties.title='Disentangling Representation Effects and Normal Reference Matching in Few-Shot Industrial Anomaly Localization'
d.core_properties.subject='English manuscript based on the current representation and matching study'
d.core_properties.author='Yuening Li'
d.core_properties.created=datetime(2026,9,23,tzinfo=timezone.utc)
d.core_properties.modified=datetime(2026,9,23,tzinfo=timezone.utc)
d.core_properties.keywords='few-shot anomaly localization; frozen visual encoders; reference matching'

def mr(t,roman=False,bold=False):
    x=OxmlElement('m:r');pr=OxmlElement('m:rPr');st=OxmlElement('m:sty');st.set(qn('m:val'),'b' if roman and bold else 'p' if roman else 'bi' if bold else 'i');pr.append(st);x.append(pr)
    wp=OxmlElement('w:rPr');f=OxmlElement('w:rFonts')
    for a in ['ascii','hAnsi']:f.set(qn('w:'+a),'Cambria Math')
    wp.append(f);x.append(wp);tt=OxmlElement('m:t');tt.text=t;x.append(tt);return x
def obj(tag,**parts):
    e=OxmlElement('m:'+tag)
    for name,children in parts.items():
        part=OxmlElement('m:'+name)
        for ch in children if isinstance(children,list) else [children]:part.append(ch)
        e.append(part)
    return e
def sub(base,index,roman=False,bold=False):
    return obj('sSub',e=[mr(base,roman,bold)],sub=index if isinstance(index,list) else [mr(index)])
def sup(base,index,roman=False,bold=False):return obj('sSup',e=[mr(base,roman,bold)],sup=[mr(index)])
def sub_sup(base,index,upper,bold=False):return obj('sSubSup',e=[mr(base,False,bold)],sub=[mr(index)],sup=[mr(upper)])
def labelindex(t):
    return [mr(part,part in ['TRI','BAL','DUP','A1','J','L','S','D','img','vis']) for part in re.split(r'(TRI|BAL|DUP|A1|img|vis|J|L|S|D)',t) if part]
def sym(t):
    t=t.replace(r'\mathrm{img}','img').replace(r'\mathrm{vis}','vis')
    if t==r'\mathcal{R}_c':return [sub('ℛ','c',True)]
    if t=='N_cd':return [sub('N','c'),mr('d')]
    if t=='nN_cd':return [mr('n'),sub('N','c'),mr('d')]
    if t==r'\tau_{vis}':return [sub('τ',[mr('vis',True)])]
    m=re.fullmatch(r'([A-Za-z]+)_\{([^}]+)\}',t) or re.fullmatch(r'([A-Za-z]+)_([A-Za-z]+)',t)
    if t=='x_i^c':return [sub_sup('x','i','c',True)]
    if m:
        b,ix=m.groups();return [sub(b,labelindex(ix),b in ['TRI','BAL','DUP','A1','R'],b in ['F','g','A','a','M'])]
    return [mr(t,t in ['J','L'] and False or t in ['P'] and False,t in ['x','F','g','A','a','M'])]
def mathrun(p,items):
    om=OxmlElement('m:oMath')
    for x in items:om.append(x)
    p._p.append(om)
def limit(name,idx):return obj('limLow',e=[mr(name,True)],lim=idx)
def summation(items):
    e=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');c=OxmlElement('m:chr');c.set(qn('m:val'),'∑');pr.append(c)
    h=OxmlElement('m:supHide');h.set(qn('m:val'),'1');pr.append(h);e.append(pr)
    e.append(obj('sub',e=[])) if False else None
    for name,children in [('sub',[mr('b')]),('sup',[]),('e',items)]:
        pp=OxmlElement('m:'+name)
        for ch in children:pp.append(ch)
        e.append(pp)
    return e
def configuration(name,mode):return sub(name,labelindex(mode),True)
def performance(name,mode):return [mr('P'),mr('(',True),configuration(name,mode),mr(')',True)]
def effect(name,mode):return sub('E',labelindex(name+','+mode))
def parg(name):return [mr('(',True),mr(name),mr(')',True)]
def distance():return [sub('d','b'),mr('(',True),mr('p'),mr(',',True),mr('r'),mr(')',True)]
def eq(n):
    if n==1:
        nd=[sub('𝒳','c',True),mr(' = {',True),sub_sup('x','i','c',True),mr(' : ',True),mr('i'),mr(' = 1, …, ',True),mr('K'),mr('}',True)]
    elif n==2:
        g1=obj('sSup',e=[sub('g',labelindex('b,p'),False,True)],sup=[mr('T',True)])
        nd=distance()+[mr(' = 1 − ',True),g1,sub('g',labelindex('b,r'),False,True)]
    elif n in [3,4]:
        mi=lambda:limit('min',[mr('r'),mr(' ∈ ',True),sub('ℛ','c',True)])
        term=[sub('w','b')]+distance()
        nd=[mr('J' if n==3 else 'L')]+parg('p')+[mr(' = ',True)]
        nd += [mi(),summation(term)] if n==3 else [summation([sub('w','b'),mi()]+distance())]
    elif n==5:nd=[mr('G')]+parg('p')+[mr(' = ',True),mr('J')]+parg('p')+[mr(' − ',True),mr('L')]+parg('p')+[mr(' ≥ 0',True)]
    elif n in [6,7]:
        name,control=('TRI','DUP') if n==6 else ('BAL','A1')
        nd=[effect(name,'t'),mr(' = ',True)]+performance(name,'t')+[mr(' − ',True)]+performance(control,'t')
    elif n in [8,9]:
        name='TRI' if n==8 else 'BAL';nd=[sub('I',[mr(name,True)]),mr(' = ',True),effect(name,'L'),mr(' − ',True),effect(name,'J')]
    elif n==10:nd=[sub('ΔI','q'),mr(' = ',True),sub('I',labelindex('q,D')),mr(' − ',True),sub('I',labelindex('q,S')),mr(',    ',True),mr('q'),mr(' ∈ {TRI, BAL}',True)]
    elif n==11:
        nd=[sub('A','t',False,True),mr(' = ',True),sub('Gauss',[mr('σ'),mr('=4',True)],True),mr('(',True),sub('Resize',[mr('H'),mr('×',True),mr('W')],True),mr('(',True),sub('a','t',False,True),mr(')),   ',True),sub('s',labelindex('img,t')),mr(' = ',True),limit('max',[mr('u')]),sub('A',labelindex('t,u'),False,True)]
    elif n==12:nd=[sub('M',labelindex('vis,t'),False,True)]+parg('u')+[mr(' = ',True),mr('1',True),mr('[',True),sub('A',labelindex('t,u'),False,True),mr(' ≥ ',True),sub('τ',[mr('vis',True)]),mr(']',True)]
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');mathrun(p,nd);p.add_run('\t('+str(n)+')')

refs=json.loads((TMP/'references.json').read_text(encoding='utf-8')) if (TMP/'references.json').exists() else {}
if isinstance(refs,list):refs={r['key']:r for r in refs}
elif 'references' in refs:refs={r['key']:r for r in refs['references']}
numbers={}
def cite(m):
    keys=[x.strip().lstrip('@') for x in m.group(1).split(';')]
    for k in keys:
        if k not in numbers:numbers[k]=len(numbers)+1
    return '['+', '.join(str(n) for n in sorted({numbers[k] for k in keys}))+']'
def inline(p,text,sub_vars=True):
    text=re.sub(r'\[@([^\]]+)\]',cite,text)
    # Match standalone mathematical labels even in tables and figure captions.
    # Exclude occurrences glued to a hyphen so that names such as K-NG or ViT-L/14 stay plain text.
    # Reference entries keep their literal text and never receive this substitution.
    if sub_vars:
        text=''.join(s if i%2 else re.sub(r'(?<![\w\-])([KHW])(?![\w\-])',r'$\1$',s) for i,s in enumerate(re.split(r'(\$[^$]+\$)',text)))
    for part in re.split(r'(\$[^$]+\$|\*\*.*?\*\*)',text):
        if part.startswith('$') and part.endswith('$'):mathrun(p,sym(part[1:-1]))
        elif part.startswith('**') and part.endswith('**'):p.add_run(part[2:-2]).bold=True
        else:p.add_run(part)
    return p
def para(text,style=None):return inline(d.add_paragraph(style=style),text)

tables=json.loads((TMP/'tables.json').read_text(encoding='utf-8'))
figures=json.loads((TMP/'figures.json').read_text(encoding='utf-8'))
table_no=0;figure_no=0
def table(key):
    global table_no
    spec=tables[key]
    if 'label' not in spec:table_no+=1
    label=spec.get('label',str(table_no))
    cap=para(f'Table {label}. '+spec['caption'],'Caption');cap.paragraph_format.keep_with_next=True
    rows=[spec['headers']]+spec['rows']
    # Group repeated values (for example one dataset name per block) by showing the value once.
    prev={};disp=[]
    for ri,row in enumerate(rows):
        r2=list(row)
        for ci in spec.get('dedupe_cols',[]):
            if str(r2[ci])==prev.get(ci,'\x00'):r2[ci]=''
            else:prev[ci]=str(r2[ci])
        disp.append(r2)
    t=d.add_table(rows=0,cols=len(disp[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=spec['widths']
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for ri,row in enumerate(disp):
        rr=t.add_row();rp=rr._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if ri==0:rp.append(OxmlElement('w:tblHeader'))
        for ci,(c,txt) in enumerate(zip(rr.cells,row)):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                ee=OxmlElement('w:'+edge);on=(edge=='top' and ri==0) or (edge=='bottom' and ri in [0,len(rows)-1]);ee.set(qn('w:val'),'single' if on else 'nil');ee.set(qn('w:sz'),'8' if ri==len(rows)-1 else '6');ee.set(qn('w:color'),'000000');bd.append(ee)
            cp.append(bd);mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                z=OxmlElement('w:'+edge);z.set(qn('w:w'),'70');z.set(qn('w:type'),'dxa');mar.append(z)
            cp.append(mar)
            p=c.paragraphs[0];inline(p,str(txt));p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
            # Chain the rows so that the caption, the whole table and its note are never split across pages.
            p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci in spec.get('left_cols',[0]) else WD_ALIGN_PARAGRAPH.CENTER
            bold=ri==0 or ri-1 in spec.get('bold_rows',[])
            for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9.5);r.bold=bold
            if bold:
                # Math runs are not part of Paragraph.runs; bold them too so an anchor row is uniformly bold.
                for mr in p._p.iter(qn('m:r')):
                    mpr=mr.find(qn('m:rPr'));sty=mpr.find(qn('m:sty')) if mpr is not None else None
                    if sty is not None:sty.set(qn('m:val'),'b' if sty.get(qn('m:val'))=='p' else 'bi')
            for rp0 in p._p.iter(qn('w:rPr')):
                z=OxmlElement('w:sz');z.set(qn('w:val'),'19');rp0.append(z)
    # Set the inherited table default to none; direct three-line cell borders prevail.
    tb=t._tbl.tblPr;bd=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        z=OxmlElement('w:'+edge);z.set(qn('w:val'),'nil');bd.append(z)
    tb.append(bd)
    if spec.get('note'):
        p=para(spec['note'],'Caption');p.paragraph_format.space_before=Pt(4)
        for run in p.runs:run.bold=False
    else:d.add_paragraph().paragraph_format.space_after=Pt(0)
def figure(key):
    global figure_no
    sp=figures[key]
    if 'label' in sp:label='Figure '+sp['label']
    else:figure_no+=1;label='Figure '+str(figure_no)
    for i,raw in enumerate(sp.get('parts',[sp['path']])):
        path=Path(raw);path=path if path.is_absolute() else ROOT/path
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(4)
        p.add_run().add_picture(str(path),width=Cm(sp.get('width',17)))
        cap=sp.get('part_captions',[sp['caption']]*len(sp.get('parts',[])))[i] if sp.get('part_captions') else (sp['caption'] if i==0 else sp.get('continuation_caption',label+' continued.'))
        cp=para((label+'. ' if i==0 else '')+cap,'Caption')
        cp.paragraph_format.keep_with_next=False

text=(TMP/'manuscript.md').read_text(encoding='utf-8').replace('{{results}}',(TMP/'results.md').read_text(encoding='utf-8'))
for line in text.splitlines():
    if not line.strip():continue
    mm=re.fullmatch(r'\{\{(eq|table|figure):([^}]+)\}\}',line)
    if mm:
        kind,key=mm.groups()
        if kind=='eq':eq(int(key))
        elif kind=='table':table(key)
        else:figure(key)
    elif line=='{{references}}':
        for key,number in numbers.items():
            if key not in refs:raise ValueError('Missing reference '+key)
            r=refs[key]
            bib=r.get('formatted') or r.get('citation') or r.get('reference') or f"{r['authors']}. {r['title']}. {r.get('venue','')}, {r['year']}."
            url=r.get('url') or r.get('primary_url')
            # Keep the approved reference format: literal citation text followed by the plain source URL.
            p=d.add_paragraph();inline(p,f'[{number}] '+bib+((' '+url) if url else ''),sub_vars=False)
            p.paragraph_format.left_indent=Cm(.65);p.paragraph_format.first_line_indent=Cm(-.65);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
            for run in p.runs:run.font.size=Pt(9.5)
    elif line.startswith('# '):
        p=para(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('#### '):para(line[5:],'Heading 3')
    elif line.startswith('### '):para(line[4:],'Heading 2')
    elif line.startswith('## '):para(line[3:],'Heading 1')
    else:para(line)

# Prevent a lead-in sentence from being separated from its display equation.
paragraphs=d.paragraphs
for i,p0 in enumerate(paragraphs[:-1]):
    nxt=paragraphs[i+1]
    if nxt._p.xpath('.//m:oMath') and re.search(r'\(\d+\)$',nxt.text) and '\t' in nxt.text:
        p0.paragraph_format.keep_with_next=True
dest=OUT/'Reference_Matching_Complete_English_20260923.docx' 
d.save(dest)
# Fix package timestamps; this affects packaging only, never numerical content.
with zipfile.ZipFile(dest) as zin:
    members={name:zin.read(name) for name in zin.namelist()}
# Drop image parts that no part of the rebuilt package references. The layout master carries
# eight unused template images (word/media/image1..8.png) whose relationships survive the
# body reset; they are not referenced by any w:drawing, so they are removed here. Numerical
# content and the 27 embedded figures are unaffected.
rels_name='word/_rels/document.xml.rels'
pruned_media=0
if rels_name in members:
    referenced={m.group(1).decode() for m in re.finditer(rb'r:(?:embed|link|id|pict|dm|lo|qs|cs|href)="([^"]+)"',members.get('word/document.xml',b''))}
    rels_data=members[rels_name]
    for m in list(re.finditer(rb'<Relationship\b[^>]*/>',rels_data)):
        chunk=m.group(0)
        if b'/image"' not in chunk:continue
        hit=re.search(rb'Target="media/([^"]+)"',chunk)
        rid=re.search(rb'Id="([^"]+)"',chunk)
        if hit is None or rid is None or rid.group(1).decode() in referenced:continue
        rels_data=rels_data.replace(chunk,b'',1)
        members.pop('word/media/'+hit.group(1).decode(),None)
        pruned_media+=1
    members[rels_name]=rels_data
with zipfile.ZipFile(dest,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as zout:
    for name in sorted(members):
        zi=zipfile.ZipInfo(name,date_time=(2026,9,23,0,0,0))
        zi.compress_type=zipfile.ZIP_DEFLATED
        zout.writestr(zi,members[name])
# Record preserved package structures; body, metadata and image relations are editable.
with zipfile.ZipFile(REF) as a,zipfile.ZipFile(dest) as b:
    preserved=[n for n in a.namelist() if n in b.namelist() and (n.startswith('word/footer') or n.startswith('word/header') or n in ['word/numbering.xml','word/theme/theme1.xml'])]
    fidelity={n:a.read(n)==b.read(n) for n in preserved}
    source_section=Document(REF).sections[0];out_section=d.sections[0]
    page={k:getattr(source_section,k)==getattr(out_section,k) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']}
assert all(fidelity.values()),fidelity
assert all(page.values()),page
assert SOURCE_SHA==hashlib.sha256(REF.read_bytes()).hexdigest()
resolved=re.sub(r'\[@([^\]]+)\]',lambda m:'['+', '.join(str(n) for n in sorted({numbers[k.strip().lstrip('@')] for k in m[1].split(';')}))+']',text)
equations={1:r'\mathcal{X}_c=\{\boldsymbol{x}_i^c:i=1,\ldots,K\}',2:r'd_b(p,r)=1-\boldsymbol{g}_{b,p}^{\mathsf{T}}\boldsymbol{g}_{b,r}',3:r'J(p)=\min_{r\in\mathcal{R}_c}\sum_b w_b d_b(p,r)',4:r'L(p)=\sum_b w_b\min_{r\in\mathcal{R}_c}d_b(p,r)',5:r'G(p)=J(p)-L(p)\geq0',6:r'E_{\mathrm{TRI},t}=P(\mathrm{TRI}_t)-P(\mathrm{DUP}_t)',7:r'E_{\mathrm{BAL},t}=P(\mathrm{BAL}_t)-P(\mathrm{A1}_t)',8:r'I_{\mathrm{TRI}}=E_{\mathrm{TRI},\mathrm{L}}-E_{\mathrm{TRI},\mathrm{J}}',9:r'I_{\mathrm{BAL}}=E_{\mathrm{BAL},\mathrm{L}}-E_{\mathrm{BAL},\mathrm{J}}',10:r'\Delta I_q=I_{q,\mathrm{D}}-I_{q,\mathrm{S}},\quad q\in\{\mathrm{TRI},\mathrm{BAL}\}',11:r'\boldsymbol{A}_t=\operatorname{Gauss}_{\sigma=4}(\operatorname{Resize}_{H\times W}(\boldsymbol{a}_t)),\quad s_{\mathrm{img},t}=\max_u \boldsymbol{A}_{t,u}',12:r'\boldsymbol{M}_{\mathrm{vis},t}(u)=\mathbf{1}[\boldsymbol{A}_{t,u}\geq\tau_{\mathrm{vis}}]'}
for num,latex in equations.items():resolved=resolved.replace('{{eq:'+str(num)+'}}','$$\n'+latex+'\\tag{'+str(num)+'}\n$$')
tn=0
for match in list(re.finditer(r'\{\{table:([^}]+)\}\}',resolved)):
    sp=tables[match[1]]
    if 'label' not in sp:tn+=1
    mt='Table '+sp.get('label',str(tn))+'. '+sp['caption']+'\n\n'
    mt+='| '+' | '.join(sp['headers'])+' |\n| '+' | '.join(['---']*len(sp['headers']))+' |\n'
    for row in sp['rows']:mt+='| '+' | '.join(map(str,row))+' |\n'
    mt+='\n'+sp.get('note','')
    resolved=resolved.replace(match[0],mt)
biblio=[]
for key,num in numbers.items():
    rr=refs[key];bib=rr.get('formatted') or rr.get('citation') or rr.get('reference') or f"{rr['authors']}. {rr['title']}. {rr.get('venue','')}, {rr['year']}."
    biblio.append(f'[{num}] '+bib+' '+(rr.get('url') or rr.get('primary_url') or ''))
resolved=resolved.replace('{{references}}','\n\n'.join(biblio))
fn=0
for match in list(re.finditer(r'\{\{figure:([^}]+)\}\}',resolved)):
    sp=figures[match[1]]
    fn+=0 if 'label' in sp else 1
    figure_label=sp.get('label',str(fn))
    panels=[]
    for j,raw in enumerate(sp.get('parts',[sp['path']])):
        cap=sp['part_captions'][j] if sp.get('part_captions') else (sp['caption'] if j==0 else sp.get('continuation_caption','Continued.'))
        panels.append('![Figure '+figure_label+' part '+str(j+1)+']('+str(raw).replace('\\','/')+')\n\n'+('Figure '+figure_label+'. ' if j==0 else '')+cap)
    resolved=resolved.replace(match[0],'\n\n'.join(panels))

assert '{{' not in resolved
(OUT/'English_Manuscript_Source.md').write_text(resolved,encoding='utf-8')
(TMP/'build_validation.json').write_text(json.dumps({'source_sha256':SOURCE_SHA,'preserved_parts':fidelity,'page_fidelity':page,'tables':table_no,'figures':figure_no,'display_equations':12,'native_math_objects':len(d.element.xpath('//m:oMath')),'references':numbers,'pruned_media_parts':pruned_media,'words_approx':len(re.findall(r"\b[\w'-]+\b",text))},indent=2),encoding='utf-8')
print(dest)
