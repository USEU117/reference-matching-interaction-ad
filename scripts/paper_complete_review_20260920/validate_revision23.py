from pathlib import Path
import json,re,hashlib,zipfile,posixpath,subprocess
from lxml import etree as E
from docx import Document
from PIL import Image,ImageDraw
import numpy as np
R=Path(__file__).resolve().parents[2];S=Path(__file__).resolve().parent;O=R/'docs/paper_complete_review_20260920';T=R/'.tmp_revision_20260923'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
checks={}
def check(key,value):
 checks[key]=bool(value)
 assert value,key
base=R/'experiments/dynamic_fusion/representation_matching_interaction_20260914'
check('frozen_common_csv',sha(base/'05_baselines_multi_dataset/baseline_common_region.csv').upper()=='3C83AB004420A4F836102CABC5F8248DEBFEBC742D8E9602FED0881823A0B8BB')
check('frozen_extension_csv',sha(base/'05_baselines_ext_20260921/baseline_common_region_ext.csv').upper()=='1C77012971A4C2EBA52512A8D7850C0DA072B8107FFFE316A74E3C39DF73EC4B')
check('frozen_template',sha(R/'docs/manuscript_polished_20260919/Reference_Matching_English_Polished_20260919.docx').upper()=='9DB99E60CD3024D1D49429641EDD6E49777BF014A3F4B7FB674C9C20338FB837')
check('no_experiment_data_changes',not subprocess.check_output(['git','status','--porcelain','--','experiments','data'],cwd=R,text=True).strip())
tables=json.loads((S/'tables.json').read_text(encoding='utf-8'))
prior=json.loads(subprocess.check_output(['git','show','HEAD:scripts/paper_complete_review_20260920/tables.json'],cwd=R,text=True,encoding='utf-8'))
check('Table11_all_fields_unchanged',tables['baselines']==prior['baselines'])
paper=O/'Reference_Matching_Complete_English_20260923.docx';deck=O/'All_Figures_Complete_20260923.pptx'
d=Document(paper);text='\n'.join(p.text for p in d.paragraphs)
check('23_tables',len(d.tables)==23);check('27_embedded_figures',len(d.inline_shapes)==27)
check('12_numbered_equations',len([p for p in d.paragraphs if re.search(r'\(\d+\)$',p.text) and '\t' in p.text])==12)
check('author',d.core_properties.author=='Yuening Li' and 'Yuening Li' in text)
check('four_result_groups',sum(p.style.name=='Heading 3' and re.match(r'4\.2\.[1-4] ',p.text) is not None for p in d.paragraphs)==4)
check('independent_discussion',text.index('5 Discussion')<text.index('6 Conclusion'))
check('related_work_no_subsections',not re.search(r'^2\.[123] ',text,re.M))
check('body_font',d.styles['Normal'].font.name=='Times New Roman' and d.styles['Normal'].font.size.pt==11)
check('figure_widths',all(abs(sh.width/360000-17)<.001 for sh in d.inline_shapes))
check('full_names',sum(t.cell(0,1).text=='Full name' for t in d.tables)==2)
check('remaining_metadata_only',set(re.findall(r'\[\[([^]]+)\]\]',text))=={'AFFILIATIONS','CORRESPONDING_AUTHOR','FUNDING','COMPETING_INTERESTS_TO_BE_CONFIRMED'})
check('no_literal_template_slots','{{' not in text)
check('no_machine_paths','My_github' not in text and 'C:\\Users' not in text)
check('no_local_matching','local matching' not in text.lower())
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
idx=json.loads((O/'FIGURE_SLIDE_INDEX.json').read_text(encoding='utf-8'))
check('63_slide_index',len(idx)==63)
native=[];media=[];styles=[]
with zipfile.ZipFile(deck) as z:
 for item in idx:
  n=item['slide'];xml=E.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
  if item['native'] if 'native' in item else False:
   native.append(n);check(f'native_slide_{n}',len(xml.findall('.//p:sp',ns))>10)
   if n==2:
    for run in xml.findall('.//a:r',ns):
     if run.find('a:t',ns).text=='c':styles.append(run.find('a:rPr',ns).get('i'))
  else:
   rels=E.fromstring(z.read(f'ppt/slides/_rels/slide{n}.xml.rels'))
   rid=xml.find('.//a:blip',ns).get('{'+ns['r']+'}embed')
   target=next(x.get('Target') for x in rels if x.get('Id')==rid)
   member=posixpath.normpath(posixpath.join('ppt/slides',target))
   match=hashlib.sha256(z.read(member)).hexdigest()==sha(R/item['image']);media.append(match)
check('all_59_PPT_image_bytes_match_sources',len(media)==59 and all(media))
check('native_slide_positions',native==[1,2,3,15])
check('italic_category_c',styles and all(v=='1' for v in styles))
with zipfile.ZipFile(paper) as z:
 doc_images={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
 specs=json.loads((S/'figures.json').read_text(encoding='utf-8'))
 check('all_27_Word_image_bytes_match_sources',all(sha(R/p) in doc_images for f in specs.values() for p in f.get('parts',[f['path']])))
coverage=[]
for f,y,h in [('fig2_matching.png',816,232),('fig3_constructions.png',762,286)]:
 im=np.asarray(Image.open(O/'figures'/f).convert('RGB'));crop=im[int(y/1060*im.shape[0]):int((y+h)/1060*im.shape[0])]
 coverage.append({'figure':f,'c_height_fraction':h/1060,'c_ink_fraction_rgb_below_140':float((crop.min(axis=2)<140).mean())})
check('compact_c_panels',coverage[0]['c_height_fraction']<=.22 and coverage[1]['c_height_fraction']<=.28 and all(x['c_ink_fraction_rgb_below_140']>=.02 for x in coverage))
stats=json.loads((T/'word_review.json').read_text(encoding='utf-8-sig'))
check('all_tables_single_page',all(t['start']==t['end'] for t in stats['tablePages']))
out={'checks':checks,'word':stats,'native_math_objects':len(d.element.xpath('//m:oMath')),'panel_c_measurements':coverage,'docx_sha256':sha(paper),'pptx_sha256':sha(deck),'native_slides':native,'ppt_bitmap_slides':len(media),'remaining_metadata':sorted(set(re.findall(r'\[\[([^]]+)\]\]',text)))}
(O/'REVISION_VALIDATION_20260923.json').write_text(json.dumps(out,indent=2),encoding='utf-8')
for start in range(28,64,6):
 sheet=Image.new('RGB',(1800,1040),'#ddd');draw=ImageDraw.Draw(sheet)
 for j,n in enumerate(range(start,min(start+6,64))):
  im=Image.open(T/f'slides/slide-{n}.png');im.thumbnail((595,493));x=(j%3)*600;y=(j//3)*520
  sheet.paste(im,(x,y+24));draw.text((x+5,y+3),f'Slide {n}',fill='black')
 sheet.save(T/f'slides/appendix-{start}.jpg')
print(json.dumps({'checks_passed':len(checks),'word_pages':stats['pages'],'native':native,'panel_c':coverage},indent=2))
